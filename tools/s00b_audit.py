"""Read-only S00B audit of user-confirmed local competition files.

Only Attachment 2 pickle files are deserialized. Attachment 3 and 4 are
enumerated with filesystem metadata; their contents are never opened.
"""
from __future__ import annotations

import argparse
import collections
import csv
import gc
import hashlib
import json
import pickle
import re
import subprocess
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {"id", "raw_text", "text", "text_bert", "audio", "vision", "annotations", "classification_labels", "regression_labels", "audio_lengths", "vision_lengths"}
SPLITS = ("train", "valid", "test")
MODALITIES = ("text", "audio", "vision")
BINS = ((1, 1, "1"), (2, 2, "2"), (3, 3, "3"), (4, 4, "4"), (5, 5, "5"), (6, 10, "6-10"), (11, 20, "11-20"), (21, 10**9, ">20"))


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, default=json_default) + "\n", encoding="utf-8")


def json_default(value: object):
    if isinstance(value, np.generic):
        return value.item()
    raise TypeError(type(value).__name__)


def file_fingerprint(path: Path) -> dict:
    stat = path.stat()
    return {"file_name": path.name, "size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "sha256": sha256(path)}


def parse_id(value: object) -> tuple[str, str] | None:
    if not isinstance(value, str) or value.count("$_$") != 1:
        return None
    left, right = value.split("$_$", 1)
    return (left, right) if left and right else None


def zero_rows(values: np.ndarray) -> np.ndarray:
    return np.all(values == 0, axis=-1)


def zero_runs(mask: np.ndarray) -> list[tuple[int, int, str]]:
    mask = np.asarray(mask, dtype=bool)
    starts = np.flatnonzero(mask & ~np.r_[False, mask[:-1]])
    ends = np.flatnonzero(mask & ~np.r_[mask[1:], False]) + 1
    return [(int(a), int(b), "whole" if a == 0 and b == len(mask) else "prefix" if a == 0 else "suffix" if b == len(mask) else "internal") for a, b in zip(starts, ends)]


def zero_structure(mask: np.ndarray) -> str:
    runs = zero_runs(mask)
    if not runs:
        return "NO_ZERO"
    if len(runs) == 1 and runs[0][2] == "whole":
        return "ALL_ZERO"
    kinds = [run[2] for run in runs]
    internal = kinds.count("internal")
    if internal > 1:
        return "MULTIPLE_INTERNAL_RUNS"
    if internal:
        return "INTERNAL_ZERO_RUN"
    if "prefix" in kinds and "suffix" in kinds:
        return "PREFIX_AND_SUFFIX"
    return "PREFIX_ZERO_ONLY" if "prefix" in kinds else "SUFFIX_ZERO_ONLY"


def run_bin(length: int) -> str:
    return next(label for low, high, label in BINS if low <= length <= high)


def length_consistency(mask: np.ndarray, lengths: np.ndarray) -> dict:
    n, t = mask.shape
    lengths = np.asarray(lengths)
    integer_like = bool(np.all(np.isfinite(lengths)) and np.all(lengths == np.floor(lengths)))
    legal = bool(integer_like and np.all((lengths >= 0) & (lengths <= t)))
    result = {"integer_like": integer_like, "legal_0_to_T": legal, "sample_count": n, "T": t}
    if not legal:
        return result
    inside = np.arange(t)[None, :] < lengths.astype(int)[:, None]
    after = ~inside
    leading = np.zeros_like(mask)
    internal = np.zeros_like(mask)
    for i in range(n):
        for a, b, kind in zero_runs(mask[i, : int(lengths[i])]):
            if kind == "prefix":
                leading[i, a:b] = True
            if kind == "internal":
                internal[i, a:b] = True
    result.update({
        "tail_after_length_all_zero_rate": float(np.mean(np.all(mask | inside, axis=1))),
        "nonzero_after_length_count": int(np.count_nonzero(~mask & after)),
        "zero_inside_declared_length_count": int(np.count_nonzero(mask & inside)),
        "leading_zero_inside_declared_length_count": int(np.count_nonzero(leading)),
        "internal_zero_inside_declared_length_count": int(np.count_nonzero(internal)),
    })
    return result


def field_schema(value: object) -> dict:
    array = np.asarray(value)
    return {"python_type": type(value).__name__, "dtype": str(array.dtype), "shape": list(array.shape), "ndim": int(array.ndim)}


def schema(data: dict, fingerprint: dict) -> tuple[dict, dict]:
    result = {**fingerprint, "python_type": type(data).__name__, "top_level_keys": list(data), "splits": {}}
    consistency = {}
    for split, item in data.items():
        n = len(item["id"]) if "id" in item else None
        fields = {key: field_schema(value) for key, value in item.items()}
        bad = [key for key, meta in fields.items() if meta["ndim"] == 0 or meta["shape"][0] != n]
        result["splits"][split] = {"sample_count": n, "field_names": list(item), "fields": fields, "EXTRA_FIELDS": sorted(set(item) - EXPECTED)}
        consistency[split] = {"sample_count": n, "mismatched_first_dimension_fields": bad, "pass": not bad}
    result["total_sample_count"] = sum(x["sample_count"] for x in result["splits"].values())
    result["specified_expected_total"] = 4850
    result["COMPETITION_SPEC_MISMATCH"] = result["total_sample_count"] != 4850
    return result, consistency


def id_audit(data: dict) -> tuple[dict, dict]:
    result = {}
    sets = {}
    for split in SPLITS:
        ids = list(data[split]["id"])
        parsed = [parse_id(x) for x in ids]
        sets[split] = set(ids)
        result[split] = {
            "sample_count": len(ids), "empty_count": sum(not x for x in ids),
            "duplicate_count": len(ids) - len(sets[split]), "malformed_count": sum(x is None for x in parsed),
            "unique_video_id_count": len({x[0] for x in parsed if x}),
        }
    for a, b in (("train", "valid"), ("train", "test"), ("valid", "test")):
        result[f"{a}_{b}_exact_id_overlap_count"] = len(sets[a] & sets[b])
        av = {x[0] for x in (parse_id(y) for y in sets[a]) if x}
        bv = {x[0] for x in (parse_id(y) for y in sets[b]) if x}
        result[f"{a}_{b}_video_id_overlap_count"] = len(av & bv)
    return result, sets


def finite_audit(data: dict) -> dict:
    result = {}
    for split in SPLITS:
        if split == "test":
            count = sum(int(np.count_nonzero(~np.isfinite(data[split][m]))) for m in MODALITIES)
            result[split] = {"nonfinite_exists": count > 0, "nonfinite_total_count": count}
        else:
            result[split] = {}
            all_zero = {}
            for mod in MODALITIES:
                a = data[split][mod]
                result[split][mod] = {"nan_count": int(np.count_nonzero(np.isnan(a))), "posinf_count": int(np.count_nonzero(np.isposinf(a))), "neginf_count": int(np.count_nonzero(np.isneginf(a)))}
                all_zero[mod] = np.all(a == 0, axis=(1, 2))
                result[split][mod]["whole_modality_all_zero_sample_count"] = int(np.count_nonzero(all_zero[mod]))
            result[split]["all_three_modalities_zero_sample_count"] = int(np.count_nonzero(all_zero["text"] & all_zero["audio"] & all_zero["vision"]))
    result["interpretation"] = "DATA_QUALITY_EVIDENCE; zero is not classified as missing or padding"
    return result


def zero_audit(data: dict, version: str) -> tuple[list[dict], list[dict], dict, dict]:
    summary, positions, masks, histograms = [], [], {}, {}
    for split in ("train", "valid"):
        masks[split] = {}
        histograms[split] = {}
        for mod in MODALITIES:
            mask = zero_rows(data[split][mod])
            masks[split][mod] = mask
            histograms[split][mod] = {str(k): v for k, v in sorted(collections.Counter(mask.sum(axis=1).tolist()).items())}
            cats = collections.Counter(zero_structure(row) for row in mask)
            runs = collections.Counter()
            for row in mask:
                for a, b, kind in zero_runs(row):
                    runs[(kind, run_bin(b - a))] += 1
            base = {"version": version, "split": split, "modality": mod, "sample_count": len(mask), "total_zero_rows": int(mask.sum())}
            summary.append({**base, "record_type": "overall", "category": "", "run_location": "", "length_bin": "", "count": int(mask.sum())})
            for category, count in sorted(cats.items()):
                summary.append({**base, "record_type": "structure", "category": category, "run_location": "", "length_bin": "", "count": count})
            for (kind, bin_label), count in sorted(runs.items()):
                summary.append({**base, "record_type": "run", "category": "", "run_location": kind, "length_bin": bin_label, "count": count})
            for pos, count in enumerate(mask.sum(axis=0)):
                positions.append({"version": version, "split": split, "modality": mod, "position": pos, "zero_count": int(count), "sample_count": len(mask)})
    return summary, positions, masks, histograms


def bert_audit(data: dict) -> tuple[dict, dict]:
    report = {}
    candidates = {}
    for split in ("train", "valid"):
        a = np.asarray(data[split]["text_bert"])
        channels = []
        for i in range(a.shape[1]):
            x = a[:, i, :]
            unique = np.unique(x)
            binary = bool(np.all(np.isin(unique, [0, 1])))
            channels.append({"channel_index": i, "unique_count": len(unique), "unique_values_if_binary": unique.tolist() if binary else None, "binary_fraction": float(np.mean(np.isin(x, [0, 1]))), "nonzero_fraction": float(np.mean(x != 0)), "has_both_binary_values": bool(binary and len(unique) == 2)})
        possible = [x["channel_index"] for x in channels if x["has_both_binary_values"]]
        candidate = possible[0] if len(possible) == 1 else None
        mask = a[:, candidate, :] == 1 if candidate is not None else None
        continuity = bool(mask is not None and all((not np.any(np.diff(row.astype(int)) == 1)) for row in mask))
        if not continuity:
            candidate, mask = None, None
        candidates[split] = mask
        report[split] = {"shape": list(a.shape), "dtype": str(a.dtype), "channel_count": a.shape[1], "channels": channels, "attention_mask_candidate": "INFERRED_ATTENTION_MASK_CHANNEL" if candidate is not None else "UNKNOWN", "candidate_channel_index": candidate}
        if mask is not None:
            lens = mask.sum(axis=1)
            report[split]["candidate_active_length_summary"] = {"min": int(lens.min()), "max": int(lens.max()), "mean": float(lens.mean())}
            report[split]["first_active_position_min_max"] = [int(np.argmax(mask, axis=1).min()), int(np.argmax(mask, axis=1).max())]
            report[split]["last_active_position_min_max"] = [int((a.shape[2] - 1 - np.argmax(mask[:, ::-1], axis=1)).min()), int((a.shape[2] - 1 - np.argmax(mask[:, ::-1], axis=1)).max())]
            report[split]["active_continuity_all_samples"] = continuity
    return report, candidates


def aligned_position_audit(zero_masks: dict, candidates: dict) -> dict:
    out = {"interpretation": "Candidate mask and structural zero evidence only; no final padding or missing definition"}
    for split in ("train", "valid"):
        active = candidates[split]
        if active is None:
            out[split] = {"status": "UNKNOWN", "reason": "No high-confidence attention-mask candidate"}
            continue
        item = {"status": "INFERRED", "candidate_active_position_count": int(active.sum()), "candidate_inactive_position_count": int((~active).sum())}
        for mod in MODALITIES:
            z = zero_masks[split][mod]
            item[mod] = {"active_zero": int(np.count_nonzero(active & z)), "active_nonzero": int(np.count_nonzero(active & ~z)), "inactive_zero": int(np.count_nonzero(~active & z)), "inactive_nonzero": int(np.count_nonzero(~active & ~z))}
        for mod in ("audio", "vision"):
            count = 0
            for row, valid in zip(zero_masks[split][mod], active):
                count += sum(int(np.count_nonzero(valid[a:b])) for a, b, kind in zero_runs(row) if kind == "internal")
            item[f"{mod}_internal_zero_run_positions_overlapping_candidate_active"] = count
        item["simultaneous_three_modality_zero_within_candidate_active"] = int(np.count_nonzero(active & zero_masks[split]["text"] & zero_masks[split]["audio"] & zero_masks[split]["vision"]))
        out[split] = item
    return out


def workbook_annotations(path: Path, allowed: set[str]) -> tuple[dict[str, str], dict]:
    # Normal worksheet access lets us read ID columns first and never access
    # label/annotation cells for quarantined test rows.
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheet = book["label"]
    header = [sheet.cell(1, i).value for i in range(1, sheet.max_column + 1)]
    indices = {str(x): i + 1 for i, x in enumerate(header)}
    result = {}
    labels = {}
    for row in range(2, sheet.max_row + 1):
        key = f"{sheet.cell(row, indices['video_id']).value}$_${sheet.cell(row, indices['clip_id']).value}"
        if key in allowed:
            result[key] = str(sheet.cell(row, indices["annotation"]).value)
            labels[key] = sheet.cell(row, indices["label"]).value
    meta = {"file_name": path.name, "row_count_excluding_header": sheet.max_row - 1, "field_names": header, "train_valid_matched_count": len(result), "test_label_cells_accessed": False}
    book.close()
    return result, {"metadata": meta, "workbook_label_for_train_valid": labels}


def label_audit(data: dict, annotations: dict[str, str], workbook_labels: dict) -> dict:
    out = {"source": "Attachment 2 label.xlsx; only train/valid annotation and label cells accessed", "TEST_LABEL_DISTRIBUTION_QUARANTINED": True}
    for split in ("train", "valid"):
        d = data[split]
        y = np.asarray(d["regression_labels"])
        c = np.asarray(d["classification_labels"])
        ids = list(d["id"])
        ann = [annotations.get(x) for x in ids]
        triples = collections.Counter((str(a), str(float(b))) for a, b in zip(ann, c))
        signs = collections.Counter(("negative" if x < 0 else "positive" if x > 0 else "zero", str(a)) for x, a in zip(y, ann))
        workbook_mismatch = sum(not np.isclose(float(workbook_labels[x]), float(v)) for x, v in zip(ids, y) if x in workbook_labels)
        finite = np.isfinite(y)
        out[split] = {
            "regression": {"dtype": str(y.dtype), "shape": list(y.shape), "min": float(np.nanmin(y)), "max": float(np.nanmax(y)), "exact_zero_count": int(np.count_nonzero(y == 0)), "negative_count": int(np.count_nonzero(y < 0)), "positive_count": int(np.count_nonzero(y > 0)), "finite_count": int(np.count_nonzero(finite)), "nan_count": int(np.count_nonzero(np.isnan(y))), "inf_count": int(np.count_nonzero(np.isinf(y))), "finite_values_legal_range": bool(np.all((-3 <= y[finite]) & (y[finite] <= 3)))},
            "classification": {"dtype": str(c.dtype), "shape": list(c.shape), "unique_values_and_counts": {str(float(k)): v for k, v in collections.Counter(c.tolist()).items()}},
            "annotation_counts": dict(collections.Counter(ann)),
            "annotation_x_classification": [{"annotation": a, "classification": b, "count": n} for (a, b), n in sorted(triples.items())],
            "regression_sign_x_annotation": [{"sign": a, "annotation": b, "count": n} for (a, b), n in sorted(signs.items())],
            "unmatched_annotation_count": ann.count(None),
            "neutral_zero_mismatch_count": sum((x == 0) != (a == "Neutral") for x, a in zip(y, ann)),
            "negative_annotation_mismatch_count": sum(x < 0 and a != "Negative" for x, a in zip(y, ann)),
            "positive_annotation_mismatch_count": sum(x > 0 and a != "Positive" for x, a in zip(y, ann)),
            "workbook_regression_mismatch_count": workbook_mismatch,
        }
    return out


def test_label_quarantine(data: dict) -> dict:
    d = data["test"]
    fields = {}
    for name in ("classification_labels", "regression_labels", "annotations"):
        if name not in d:
            fields[name] = {"exists": False}
            continue
        a = np.asarray(d[name])
        info = {"exists": True, "dtype": str(a.dtype), "shape": list(a.shape)}
        if np.issubdtype(a.dtype, np.number):
            info["finite"] = bool(np.all(np.isfinite(a)))
            info["legal_range"] = bool(np.all((-3 <= a) & (a <= 3))) if name == "regression_labels" else bool(np.all(np.isin(a, [0, 1, 2]))) if name == "classification_labels" else None
        fields[name] = info
    return {"TEST_LABEL_DISTRIBUTION_QUARANTINED": True, "fields": fields}


def capture_cross_version(data: dict) -> dict:
    out = {}
    for split in SPLITS:
        d = data[split]
        out[split] = {"ids": list(d["id"]), "classification_labels": np.asarray(d["classification_labels"]).copy(), "regression_labels": np.asarray(d["regression_labels"]).copy()}
        if split != "test":
            out[split].update({"raw_text": list(d["raw_text"]), "text": np.asarray(d["text"]).copy()})
    return out


def compare_versions(left: dict, right: dict) -> dict:
    out = {}
    for split in SPLITS:
        a, b = left[split], right[split]
        same_set = set(a["ids"]) == set(b["ids"])
        same_order = a["ids"] == b["ids"]
        item = {"sample_count_equal": len(a["ids"]) == len(b["ids"]), "id_set_equal": same_set, "id_order_equal": same_order}
        if same_set:
            order = [b["ids"].index(k) for k in a["ids"]] if not same_order else slice(None)
            item["labels_equal"] = bool(np.array_equal(a["classification_labels"], b["classification_labels"][order]) and np.array_equal(a["regression_labels"], b["regression_labels"][order]))
            if split != "test":
                item["raw_text_equal"] = all(x == y for x, y in zip(a["raw_text"], np.asarray(b["raw_text"], dtype=object)[order]))
                item["raw_text_aggregate_sha256_equal"] = hashlib.sha256("\n".join(a["raw_text"]).encode()).hexdigest() == hashlib.sha256("\n".join(np.asarray(b["raw_text"], dtype=object)[order]).encode()).hexdigest()
                item["text_exact_equal"] = bool(np.array_equal(a["text"], b["text"][order]))
                item["text_max_abs_diff"] = float(np.max(np.abs(a["text"] - b["text"][order])))
        out[split] = item
    return out


def attachment1_inventory(path: Path) -> dict:
    videos = list(path.rglob("*.mp4"))
    sheets = list(path.rglob("label-100.xlsx"))
    folders = {p.parent.name for p in videos}
    result = {"video_file_count": len(videos), "video_id_folder_count": len(folders), "label_100_exists": len(sheets) == 1, "specified_video_count": 100, "specified_video_id_folder_count": 37, "video_count_match": len(videos) == 100, "folder_count_match": len(folders) == 37}
    if len(sheets) == 1:
        book = openpyxl.load_workbook(sheets[0], read_only=False, data_only=True)
        sheet = book["label"]
        header = [sheet.cell(1, i).value for i in range(1, sheet.max_column + 1)]
        indices = {str(x): i + 1 for i, x in enumerate(header)}
        video_keys = {(p.parent.name, p.stem) for p in videos}
        label_keys = {(str(sheet.cell(i, indices["video_id"]).value), str(sheet.cell(i, indices["clip_id"]).value)) for i in range(2, sheet.max_row + 1)} if {"video_id", "clip_id"} <= set(indices) else set()
        result.update({"excel_row_count_excluding_header": sheet.max_row - 1, "required_columns_present": {x: x in indices for x in ("video_id", "clip_id", "text", "label", "annotation")}, "video_without_label_count": len(video_keys - label_keys), "label_without_video_count": len(label_keys - video_keys), "duplicate_video_key_count": len(videos) - len(video_keys), "duplicate_label_key_count": (sheet.max_row - 1) - len(label_keys)})
        book.close()
    return result


def attachment_file_inventory(path: Path, number: int) -> dict:
    # Deliberately no open/read/deserialization calls in this function.
    files = [p for p in path.rglob("*") if p.is_file()]
    suffixes = collections.Counter(p.suffix.lower() for p in files)
    def version_of(p: Path) -> str | None:
        parts = [x.name for x in (p, *p.parents) if x != path]
        if any("未对齐" in name or "unaligned" in name.lower() for name in parts):
            return "unaligned"
        if any("对齐" in name or "aligned" in name.lower() for name in parts):
            return "aligned"
        return None
    versions = {v: sum(version_of(p) == v for p in files) for v in ("aligned", "unaligned")}
    result = {"exists": path.is_dir(), "file_count": len(files), "extension_counts": dict(suffixes), "total_size_bytes": sum(p.stat().st_size for p in files), "version_filename_or_directory_counts": versions}
    if number == 3:
        result["CONTENT_NOT_INSPECTED"] = True
    else:
        video = [p for p in files if p.suffix.lower() == ".mp4"]
        feature = [p for p in files if p.suffix.lower() == ".pkl"]
        per_version = {}
        for v in ("aligned", "unaligned"):
            f = [p for p in feature if version_of(p) == v]
            videos = [p for p in video if version_of(p) == v]
            per_version[v] = {"feature_file_count": len(f), "video_file_count": len(videos), "feature_video_stem_match_count": len({p.stem for p in f} & {p.stem for p in videos})}
        result.update({"feature_file_count": len(feature), "video_count": len(video), "per_version": per_version, "FEATURE_CONTENT_NOT_INSPECTED": True, "VIDEO_CONTENT_NOT_INSPECTED": True})
    return result


def csv_write(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=ROOT / "configs" / "paths.local.json")
    parser.add_argument("--run-id")
    args = parser.parse_args()
    config = json.loads(args.config.read_text(encoding="utf-8-sig"))
    if config.get("trusted_competition_pickle") is not True:
        raise RuntimeError("Trusted official pickle must be explicitly confirmed in local config")
    data_root = Path(config["data_root"])
    if not data_root.is_dir():
        raise RuntimeError("Local official data root unavailable")
    baseline = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    run_id = args.run_id or datetime.now().strftime("%Y%m%d-%H%M%S") + "-S00B-" + baseline[:8]
    run = ROOT / "reports" / "runs" / run_id
    public = run / "public"
    private = run / "private"
    artifacts = run / "artifacts"
    for p in (public, private, artifacts, ROOT / "reports" / "data_audit"):
        p.mkdir(parents=True, exist_ok=True)
    audit = ROOT / "reports" / "data_audit"
    attach2 = data_root / "附件2-数据集特征文件"
    sources = {v: attach2 / f"{v}_50.pkl" for v in ("aligned", "unaligned")}
    before = {v: file_fingerprint(p) for v, p in sources.items()}
    write_json(public / "SOURCE_HASH_BEFORE.json", before)
    schemas, splits, ids, finite, test_labels, bert = {}, {}, {}, {}, {}, {}
    zero_summary, zero_positions, zero_hist = [], [], {}
    length_report = {}
    aligned_position_report = {}
    left_capture = None
    annotations = workbook_labels = None
    for version, source in sources.items():
        with source.open("rb") as stream:
            data = pickle.load(stream)
        schemas[version], splits[version] = schema(data, before[version])
        ids[version], idsets = id_audit(data)
        finite[version] = finite_audit(data)
        test_labels[version] = test_label_quarantine(data)
        if version == "aligned":
            allowed = idsets["train"] | idsets["valid"]
            annotations, wb = workbook_annotations(attach2 / "label.xlsx", allowed)
            workbook_labels = wb["workbook_label_for_train_valid"]
            workbook_meta = wb["metadata"]
            label_report = label_audit(data, annotations, workbook_labels)
        else:
            second_label_report = label_audit(data, annotations, workbook_labels)
        rows, positions, masks, hist = zero_audit(data, version)
        zero_summary += rows
        zero_positions += positions
        zero_hist[version] = hist
        bert[version], candidates = bert_audit(data)
        if version == "aligned":
            aligned_position_report = aligned_position_audit(masks, candidates)
            left_capture = capture_cross_version(data)
        else:
            for split in ("train", "valid"):
                length_report[split] = {}
                for mod in ("audio", "vision"):
                    key = f"{mod}_lengths"
                    length_report[split][mod] = length_consistency(masks[split][mod], np.asarray(data[split][key])) if key in data[split] else {"status": "UNKNOWN", "reason": "length field absent"}
            right_capture = capture_cross_version(data)
            version_report = compare_versions(left_capture, right_capture)
        del data, masks, candidates
        gc.collect()
    after = {v: file_fingerprint(p) for v, p in sources.items()}
    mutation = {v: before[v] == after[v] for v in sources}
    write_json(public / "SOURCE_HASH_AFTER.json", after)
    write_json(audit / "source_mutation_check.json", {"unchanged": mutation, "before": before, "after": after})
    for version in ("aligned", "unaligned"):
        write_json(audit / f"schema_{version}.json", schemas[version])
    write_json(audit / "split_consistency.json", {"versions": splits, "expected_total_per_version": 4850, "actual_totals": {v: schemas[v]["total_sample_count"] for v in schemas}, "COMPETITION_SPEC_MISMATCH": any(schemas[v]["COMPETITION_SPEC_MISMATCH"] for v in schemas)})
    write_json(audit / "id_integrity.json", ids)
    write_json(audit / "label_mapping_train_valid.json", {"workbook": workbook_meta, "aligned": label_report, "unaligned": second_label_report, "test": test_labels})
    write_json(audit / "version_consistency.json", version_report)
    csv_write(audit / "zero_run_summary.csv", zero_summary)
    csv_write(audit / "zero_position_summary.csv", zero_positions)
    write_json(audit / "zero_count_histograms.json", zero_hist)
    write_json(audit / "text_bert_diagnostics.json", bert)
    write_json(audit / "length_consistency_unaligned.json", length_report)
    write_json(audit / "aligned_positional_diagnostics.json", aligned_position_report)
    write_json(audit / "finite_value_audit.json", finite)
    write_json(audit / "attachment1_inventory.json", attachment1_inventory(data_root / "附件1-数据集原始多模态样本"))
    write_json(audit / "attachment3_inventory.json", attachment_file_inventory(data_root / "附件3-模态缺失特征样本", 3))
    write_json(audit / "attachment4_inventory.json", attachment_file_inventory(data_root / "附件4-可解释专项视频样本与特征文件", 4))
    write_json(run / "RUN.json", {"task_id": "S00B_REAL_DATA_AUDIT", "run_id": run_id, "baseline_sha": baseline, "data_kind": "real_official_local", "split_usage": {"train": "full audit", "valid": "full audit", "test": "quarantined metadata and legal checks only"}, "source_files": [p.name for p in sources.values()], "source_unchanged": mutation, "TEST_LABEL_DISTRIBUTION_QUARANTINED": True, "attachment3_content_inspected": False, "attachment4_feature_content_inspected": False, "attachment4_video_content_inspected": False})
    print(json.dumps({"run_id": run_id, "actual_totals": {v: schemas[v]["total_sample_count"] for v in schemas}, "source_unchanged": mutation, "annotation_match": workbook_meta["train_valid_matched_count"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
